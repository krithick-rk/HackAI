import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Rich Corporate Dark Navy Palette
    COLOR_BG = RGBColor(15, 23, 42)            # Rich Slate Navy (#0F172A)
    COLOR_CARD_BG = RGBColor(30, 41, 59)       # Elevated Navy Card (#1E293B)
    COLOR_CARD_BORDER = RGBColor(51, 65, 85)   # Border Line (#334155)
    
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252)  # Bright White (#F8FAFC)
    COLOR_TEXT_BODY = RGBColor(226, 232, 240)  # Off-White Body (#E2E8F0)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Soft Cool Gray (#94A3B8)
    
    COLOR_ACCENT_BLUE = RGBColor(56, 189, 248) # Sapphire Cyan (#38BDF8)
    COLOR_ACCENT_GOLD = RGBColor(251, 191, 36) # Executive Gold (#FBBF24)
    COLOR_SEV_HIGH = RGBColor(248, 113, 113)   # Soft Crimson Red (#F87171)
    COLOR_SEV_MED = RGBColor(251, 146, 60)     # Warm Amber (#FB923C)
    COLOR_SUCCESS = RGBColor(74, 222, 128)     # Emerald Green (#4ADE80)
    COLOR_PURPLE = RGBColor(192, 132, 252)    # Soft Purple (#C084FC)
    
    COLOR_TBL_HEADER_BG = RGBColor(30, 58, 138)# Deep Royal Blue (#1E3A8A)

    def set_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SETS CHENNAI — TEAM LOGICRON  |  HACK@DAC 2026 CALIPTRA AUDIT"):
        set_background(slide)
        
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT_BLUE

        # Main Title (Large & Bold)
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.65), Inches(12.333), Inches(0.85))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN

        # Accent Line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.45), Inches(12.333), Inches(0.02)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide (Advisor & Participants)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1)

    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    accent_bar.line.fill.background()

    add_card(slide1, 0.5, 0.5, 12.333, 6.5, bg_color=COLOR_CARD_BG, border_color=COLOR_ACCENT_BLUE)

    tb = slide1.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(6.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "HACK@DAC 2026 — PHASE 2 CALIPTRA RTL AUDIT"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE
    p.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "RTL Security Analysis & Verification Strategy"
    p2.font.size = Pt(34)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MAIN
    p2.space_after = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = "Comprehensive Audit of 29 IP Cores: Stage 0 Recon ➔ Phase 1: AI-Orchestrated Dynamic Simulation ➔ Phase 2: Python SAST & AI Validation Pipeline"
    p3.font.size = Pt(17)
    p3.font.color.rgb = COLOR_TEXT_BODY
    p3.space_after = Pt(22)

    p_team_hdr = tf.add_paragraph()
    p_team_hdr.text = "TEAM LOGICRON"
    p_team_hdr.font.size = Pt(20)
    p_team_hdr.font.bold = True
    p_team_hdr.font.color.rgb = COLOR_ACCENT_GOLD
    p_team_hdr.space_after = Pt(4)

    p_org = tf.add_paragraph()
    p_org.text = "Society for Electronic Transactions and Security (SETS), Chennai, India"
    p_org.font.size = Pt(15)
    p_org.font.bold = True
    p_org.font.color.rgb = COLOR_TEXT_MAIN
    p_org.space_after = Pt(12)

    p_adv = tf.add_paragraph()
    p_adv.text = "Advisor:  A. Suganya"
    p_adv.font.size = Pt(16)
    p_adv.font.bold = True
    p_adv.font.color.rgb = COLOR_ACCENT_BLUE
    p_adv.space_after = Pt(6)

    participants = ["Eswari Devi", "Renita .J", "Sathish Kumar", "Kushana Shyam Sundar"]
    p_part = tf.add_paragraph()
    p_part.text = "Participants:  " + "   •   ".join(participants)
    p_part.font.size = Pt(15)
    p_part.font.color.rgb = COLOR_TEXT_BODY
    p_part.space_after = Pt(22)

    p4 = tf.add_paragraph()
    p4.text = "STATUS: Phase 1 Dynamic Simulation (34 Bugs) Complete  |  Phase 2 Static SAST Pipeline Operational"
    p4.font.size = Pt(15)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_SUCCESS

    # ==========================================
    # SLIDE 2: Clean Blank Team Photo Slide
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Team LOGICRON — Society for Electronic Transactions and Security (SETS)", "TEAM PHOTO")

    # Clean Blank Card Canvas for pasting the Group Picture
    add_card(slide2, 0.5, 1.65, 12.333, 5.4, bg_color=COLOR_CARD_BG, border_color=COLOR_ACCENT_GOLD)

    # ==========================================
    # SLIDE 3: Workflow Architecture
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Security Audit Workflow & Phase Separation")

    stages = [
        ("STAGE 0: RECONNAISSANCE", "IP Discovery & Mapping", [
            "• Scope & IP Tree Mapping (29 IPs)",
            "• Register & RDL Specification Audit",
            "• Boundary & Interface Triage",
            "• Threat Modeling & Security Goals"
        ], COLOR_ACCENT_GOLD),
        ("PHASE 1: DYNAMIC ANALYSIS", "AI-Orchestrated Simulation", [
            "• Claude Code CLI Integration",
            "• Orchestrator AI Batch Scheduling",
            "• Worker AI Instances & Custom Tools",
            "• Runnable PoC Exploits & VCD Traces"
        ], COLOR_ACCENT_BLUE),
        ("PHASE 2: STATIC SAST PIPELINE", "Python Engine & AI Validation", [
            "• Modular Parallel Static Scanning",
            "• Rule-Based Vulnerability Detection",
            "• Claude AI Root-Cause & FP Reduction",
            "• Exploit Generation & Reproduction"
        ], COLOR_SUCCESS)
    ]

    for i, (stg_tag, stg_title, stg_bullets, stg_color) in enumerate(stages):
        left = 0.5 + i * 4.15
        top = 1.65
        add_card(slide3, left, top, 3.9, 5.4, border_color=stg_color)

        tb = slide3.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.25), Inches(3.4), Inches(4.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stg_tag
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = stg_color
        p.space_after = Pt(6)
        
        p_head = tf.add_paragraph()
        p_head.text = stg_title
        p_head.font.size = Pt(19)
        p_head.font.bold = True
        p_head.font.color.rgb = COLOR_TEXT_MAIN
        p_head.space_after = Pt(18)

        for b in stg_bullets:
            p_b = tf.add_paragraph()
            p_b.text = b
            p_b.font.size = Pt(15)
            p_b.font.color.rgb = COLOR_TEXT_BODY
            p_b.space_after = Pt(14)

    # ==========================================
    # SLIDE 4: Common Recon — 29 IP Inventory Table
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Stage 0 — Target Subsystem Inventory & Tier Breakdown")

    table_shape = slide4.shapes.add_table(6, 4, Inches(0.5), Inches(1.65), Inches(12.333), Inches(5.4))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(3.6)
    table.columns[2].width = Inches(5.333)
    table.columns[3].width = Inches(1.6)

    headers = ["Tier", "Subsystem Category", "Target IP Cores Included", "Module Count"]
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TBL_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE

    rows_data = [
        ("Tier 1", "Crypto Cores", "AES, SHA-256, SHA-512, SHA-512M, SHA-3, HMAC, HMAC-DRBG, KMAC, ECC, DOE, ML-DSA", "11 IPs"),
        ("Tier 2", "Storage & Access Control", "Key Vault, PCR Vault, Data Vault, Lifecycle Controller", "4 IPs"),
        ("Tier 3", "Entropy & RNG", "CSRNG, Entropy Source, EDN", "3 IPs"),
        ("Tier 4", "Interfaces & CPU Core", "SoC Interface, SPI Host, UART, VeeR EL2 RISC-V Core", "4 IPs"),
        ("Tier 5", "Bus Fabric & Infrastructure", "AXI, AHB-Lite Bus, Caliptra TL-UL, Caliptra Prim, Prim Generic, Shared Libs, Top Integration", "7 IPs")
    ]

    for row_idx, r_data in enumerate(rows_data, start=1):
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_BG
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(14)
            p.font.color.rgb = COLOR_TEXT_BODY

    # ==========================================
    # SLIDE 5: Phase 1 — Dynamic Analysis Methodology
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Phase 1 — Dynamic Analysis: Claude Code & AI Orchestration")

    pillars = [
        ("1. Claude Code CLI Integration", "Utilized Claude Code CLI integrated directly with SystemVerilog simulation tools (Verilator/Iverilog) to automate testbench generation and waveform analysis.", COLOR_ACCENT_BLUE),
        ("2. Batch-Wise AI Orchestration", "A central Orchestrator AI instance manages target IP analysis batch-wise, allocating isolated IP modules to specialized worker agents in parallel.", COLOR_ACCENT_GOLD),
        ("3. Tool-Assisted Bug Evaluation", "Worker AI instances invoke dedicated verification tools to run simulation stimulus, parse trace logs, and evaluate bugs for their assigned IP core.", COLOR_SUCCESS),
        ("4. Verified Exploits & Waveforms", "Every accepted bug requires a runnable SystemVerilog PoC exploit and a VCD/FST waveform trace proving the vulnerability with 0 false positives.", COLOR_SEV_HIGH)
    ]

    for i, (p_title, p_desc, p_color) in enumerate(pillars):
        col = i % 2
        row = i // 2
        left = 0.5 + col * 6.25
        top = 1.65 + row * 2.7
        add_card(slide5, left, top, 6.05, 2.5, border_color=COLOR_CARD_BORDER)

        tb = slide5.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.25), Inches(5.55), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = p_color
        p.space_after = Pt(10)

        p_body = tf.add_paragraph()
        p_body.text = p_desc
        p_body.font.size = Pt(15)
        p_body.font.color.rgb = COLOR_TEXT_BODY

    # ==========================================
    # SLIDE 6: Phase 1 — Dynamic Results & Metrics
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Phase 1 — Dynamic Analysis Results & Compute Ledger")

    stats = [
        ("34", "Confirmed Bugs\n(0 False Positives)", COLOR_SEV_HIGH),
        ("34", "Runnable PoC Exploits\n+ VCD Waveforms", COLOR_SUCCESS),
        ("7.88M", "Total Compute Tokens\n(30 Worker Spawns)", COLOR_ACCENT_GOLD),
        ("29 / 29", "IP Cores Tested\n(100% Coverage)", COLOR_ACCENT_BLUE)
    ]

    for i, (val, lbl, clr) in enumerate(stats):
        left = 0.5 + i * 3.1
        add_card(slide6, left, 1.65, 2.95, 1.7, border_color=clr)
        tb = slide6.shapes.add_textbox(Inches(left + 0.1), Inches(1.75), Inches(2.75), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = clr
        p_sub = tf.add_paragraph()
        p_sub.text = lbl
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

    table_shape = slide6.shapes.add_table(6, 4, Inches(0.5), Inches(3.55), Inches(12.333), Inches(3.5))
    table = table_shape.table
    table.columns[0].width = Inches(3.3)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(3.533)
    table.columns[3].width = Inches(3.0)

    headers = ["Top IP Core Name", "Tokens Consumed", "Execution / Spawns", "Dynamic Findings"]
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TBL_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE

    top_ips_data = [
        ("SHA-3", "676,655 tokens", "2 Spawns (~5.0h)", "1 Bug (AHB Bus DoS)"),
        ("Top Integration", "635,160 tokens", "1 Spawn (~8.3h)", "1 Bug (OTP Fuse Gate Bypass)"),
        ("ML-DSA / Adams Bridge", "423,462 tokens", "1 Spawn (~1.4h)", "1 Bug (Address Decode Collision)"),
        ("DOE (Deobfuscation Engine)", "381,629 tokens", "1 Spawn (~0.8h)", "1 Bug (Command Replay)"),
        ("AES Core", "355,032 tokens", "1 Spawn (~6.1h)", "1 Bug (Reseed Latch)")
    ]

    for row_idx, r_data in enumerate(top_ips_data, start=1):
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_BG
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(14)
            p.font.color.rgb = COLOR_TEXT_BODY if col_idx != 1 else COLOR_ACCENT_BLUE

    # ==========================================
    # SLIDE 7: Dynamic Deep Dive 1 — KV-002
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Dynamic Deep Dive 1 — Key Vault Key Wiping [KV-002]")

    add_card(slide7, 0.5, 1.65, 3.9, 5.4, border_color=COLOR_SEV_HIGH)
    tb = slide7.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(3.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BUG SPECIFICATION"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE
    p.space_after = Pt(14)

    meta = [
        ("Bug ID:", "KV-002"),
        ("Target File:", "src/keyvault/rtl/kv_top.sv"),
        ("CVSSv3.1 Score:", "8.4 (High Severity)"),
        ("Impacted IP:", "Key Vault (KV)"),
        ("Exploit Path:", "work/exploits/KV-002_exploit.sv"),
        ("Waveform Trace:", "work/waves/KV-002.vcd")
    ]
    for k, v in meta:
        p_k = tf.add_paragraph()
        p_k.text = k
        p_k.font.size = Pt(13)
        p_k.font.color.rgb = COLOR_TEXT_MUTED
        p_v = tf.add_paragraph()
        p_v.text = v
        p_v.font.size = Pt(15)
        p_v.font.bold = True
        p_v.font.color.rgb = COLOR_SEV_HIGH if "CVSS" in k else COLOR_TEXT_BODY
        p_v.space_after = Pt(6)

    add_card(slide7, 4.6, 1.65, 8.233, 5.4)
    tb = slide7.shapes.add_textbox(Inches(4.85), Inches(1.85), Inches(7.733), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    sections = [
        ("Vulnerability Mechanism:", "When any two write clients attempt simultaneous transfer (`kv_multi_write_err`), the error handling logic unconditionally zeroizes ALL 24 key slots in the Key Vault — including write-locked and use-locked key entries that had zero involvement in the collision."),
        ("Security Consequence:", "Complete Key Vault Denial of Service (DoS). An unprivileged write client can intentionally trigger a single collision with a low-privilege slot to instantly wipe critical root keys, HMAC keys, and unlocked secret material across the entire RoT."),
        ("Evidence & Verification:", "Proven via Verilator PoC exploit testbench (`KV-002_exploit.sv`). VCD trace shows simultaneous writes to Slot 1 & 2 causing `slot_zeroize_bus` to assert across all 24 slots, clearing locked Slot 0 instantly.")
    ]

    for title, desc in sections:
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT_BLUE
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_after = Pt(16)

    # ==========================================
    # SLIDE 8: Dynamic Deep Dive 2 — ECC-001
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Dynamic Deep Dive 2 — ECC Verification Register Leak [ECC-001]")

    add_card(slide8, 0.5, 1.65, 3.9, 5.4, border_color=COLOR_SEV_HIGH)
    tb = slide8.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(3.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BUG SPECIFICATION"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE
    p.space_after = Pt(14)

    meta = [
        ("Bug ID:", "ECC-001"),
        ("Target File:", "src/ecc/rtl/ecc_top.sv"),
        ("CVSSv3.1 Score:", "8.1 (High Severity)"),
        ("Impacted IP:", "Elliptic Curve Cryptography"),
        ("Exploit Path:", "work/exploits/ECC-001_exploit.sv"),
        ("Waveform Trace:", "work/waves/ECC-001.vcd")
    ]
    for k, v in meta:
        p_k = tf.add_paragraph()
        p_k.text = k
        p_k.font.size = Pt(13)
        p_k.font.color.rgb = COLOR_TEXT_MUTED
        p_v = tf.add_paragraph()
        p_v.text = v
        p_v.font.size = Pt(15)
        p_v.font.bold = True
        p_v.font.color.rgb = COLOR_SEV_HIGH if "CVSS" in k else COLOR_TEXT_BODY
        p_v.space_after = Pt(6)

    add_card(slide8, 4.6, 1.65, 8.233, 5.4)
    tb = slide8.shapes.add_textbox(Inches(4.85), Inches(1.85), Inches(7.733), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    sections = [
        ("Vulnerability Mechanism:", "The `ECC_VERIFY_R` output register is only cleared upon explicit `ZEROIZE` or debug entry, but is NEVER cleared when an invalid input is passed to the ECDSA signature verification engine hardware error path."),
        ("Security Consequence:", "Signature Verification Spoofing. If a prior valid signature verification succeeded (setting `ECC_VERIFY_R = VALID`), a subsequent invalid/forged signature submission that triggers an input error fails to reset `ECC_VERIFY_R`, leaving stale VALID status active for firmware."),
        ("Evidence & Verification:", "Confirmed with PoC exploit simulation (`ECC-001_exploit.sv`). Passing a malformed R/S pair causes hardware error flags to set, but `ECC_VERIFY_R` retains `0x1` (VALID) continuously.")
    ]

    for title, desc in sections:
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT_BLUE
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_after = Pt(16)

    # ==========================================
    # SLIDE 9: Dynamic Deep Dive 3 — INTEG-001
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Dynamic Deep Dive 3 — Top Integration OTP Fuse Gate [INTEG-001]")

    add_card(slide9, 0.5, 1.65, 3.9, 5.4, border_color=COLOR_SEV_HIGH)
    tb = slide9.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(3.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BUG SPECIFICATION"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE
    p.space_after = Pt(14)

    meta = [
        ("Bug ID:", "INTEG-001"),
        ("Target File:", "src/integration/rtl/caliptra_top.sv:1341"),
        ("CVSSv3.1 Score:", "7.9 (High Severity)"),
        ("Impacted IP:", "Top Integration (INTEG)"),
        ("Exploit Path:", "work/exploits/INTEG-001_exploit.sv"),
        ("Waveform Trace:", "work/waves/INTEG-001.vcd")
    ]
    for k, v in meta:
        p_k = tf.add_paragraph()
        p_k.text = k
        p_k.font.size = Pt(13)
        p_k.font.color.rgb = COLOR_TEXT_MUTED
        p_v = tf.add_paragraph()
        p_v.text = v
        p_v.font.size = Pt(15)
        p_v.font.bold = True
        p_v.font.color.rgb = COLOR_SEV_HIGH if "CVSS" in k else COLOR_TEXT_BODY
        p_v.space_after = Pt(6)

    add_card(slide9, 4.6, 1.65, 8.233, 5.4)
    tb = slide9.shapes.add_textbox(Inches(4.85), Inches(1.85), Inches(7.733), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    sections = [
        ("Vulnerability Mechanism:", "`caliptra_top.sv` hardcodes `entropy_src`'s OTP fuse permission (`otp_en_entropy_src_fw_over_i`) to constant `MuBi8True`. This collapses the intended two-factor safety gate (OTP Fuse AND SW register) into SW-only control."),
        ("Security Consequence:", "TRNG Entropy Injection Attack. Malicious firmware alone can force the TRNG entropy source into firmware-override mode in ANY lifecycle state (including Production). Firmware can inject arbitrary 384-bit attacker-chosen entropy feeding CSRNG seeds!"),
        ("Evidence & Verification:", "Differential PoC exploit verified against unmodified `caliptra_top.sv` vs properly-fused logic. The as-built core matches attacker-chosen entropy byte-for-byte in CSRNG.")
    ]

    for title, desc in sections:
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT_BLUE
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_after = Pt(16)

    # ==========================================
    # SLIDE 10: Phase 2 — SAST Validation Methodology
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Hardware RTL SAST Validation Methodology", "PHASE 2: STATIC ANALYSIS PIPELINE")

    sast_steps = [
        ("1. Python Analysis Framework", "Pass project directory to SAST framework to automatically traverse SystemVerilog source files."),
        ("2. Parallel Static Analysis", "Divide analysis into 10 independent scanning modules executing concurrently for high performance."),
        ("3. Centralized Finding Collection", "Collect findings across modules, merge into central database, and normalize info."),
        ("4. Internal SAST Rule Scan", "Execute Python SAST engine using validated security rules to correlate detection.")
    ]

    for i, (step_title, step_desc) in enumerate(sast_steps):
        left = 0.5 + i * 3.1
        top = 1.65
        add_card(slide10, left, top, 2.95, 4.1, border_color=COLOR_ACCENT_BLUE)

        tb = slide10.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.2), Inches(2.65), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = step_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_GOLD
        p.space_after = Pt(10)

        p_b = tf.add_paragraph()
        p_b.text = step_desc
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_BODY

        if i < 3:
            arrow_box = slide10.shapes.add_textbox(Inches(left + 2.85), Inches(top + 1.7), Inches(0.3), Inches(0.4))
            tf_a = arrow_box.text_frame
            p_a = tf_a.paragraphs[0]
            p_a.text = "➔"
            p_a.font.size = Pt(18)
            p_a.font.bold = True
            p_a.font.color.rgb = COLOR_ACCENT_BLUE

    # Scanning Modules Box under Workflow
    add_card(slide10, 0.5, 5.9, 12.333, 1.15, bg_color=COLOR_TBL_HEADER_BG, border_color=COLOR_ACCENT_GOLD)
    tb_mod = slide10.shapes.add_textbox(Inches(0.7), Inches(5.95), Inches(11.933), Inches(1.05))
    tf_mod = tb_mod.text_frame
    tf_mod.word_wrap = True

    p_mod_title = tf_mod.paragraphs[0]
    p_mod_title.text = "PARALLEL SCANNING MODULE CATEGORIES & COVERAGE EFFICIENCY"
    p_mod_title.font.size = Pt(12)
    p_mod_title.font.bold = True
    p_mod_title.font.color.rgb = COLOR_ACCENT_GOLD
    p_mod_title.space_after = Pt(2)

    p_mod_body = tf_mod.add_paragraph()
    p_mod_body.text = "Modules: Access Control | Secure Boot | OTP/Fuse Security | Debug Interface | Lifecycle State | Register Protection | Memory Protection | Information Disclosure | Cryptographic Misuse | Privilege Escalation"
    p_mod_body.font.size = Pt(12)
    p_mod_body.font.color.rgb = COLOR_TEXT_MAIN
    p_mod_body.space_after = Pt(4)

    p_note = tf_mod.add_paragraph()
    p_note.text = '"Multiple independent analysis modules improve coverage while reducing overall analysis time."'
    p_note.font.size = Pt(12)
    p_note.font.bold = True
    p_note.font.color.rgb = COLOR_SUCCESS

    # ==========================================
    # SLIDE 11: Phase 2 — Consolidated AI-Assisted Validation & Reporting Pipeline
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "AI-Assisted Validation & Reporting Pipeline", "PHASE 2: AI VALIDATION & REPORTING")

    ai_steps = [
        ("1. Claude AI (CLI) Analysis", [
            "• Root cause analysis",
            "• RTL code inspection",
            "• Security impact evaluation",
            "• Vulnerability validation"
        ], COLOR_ACCENT_BLUE),
        ("2. Duplicate & FP Filter", [
            "Filters findings matching:",
            "• Same RTL logic & root cause",
            "• Same attack path & module",
            "• Invalid reachability/intent"
        ], COLOR_ACCENT_GOLD),
        ("3. Exploit Generation", [
            "AI creates reproduction evidence:",
            "• Standalone PoC exploits",
            "• Simulation-based validation",
            "• Reproduction verification"
        ], COLOR_PURPLE),
        ("4. Final Security Report", [
            "Consolidated Final Deliverables:",
            "• Verified unique findings",
            "• Technical report & summary",
            "• Confirmed exploit evidence"
        ], COLOR_SUCCESS)
    ]

    for i, (a_title, a_bullets, a_color) in enumerate(ai_steps):
        left = 0.5 + i * 3.1
        top = 1.65
        add_card(slide11, left, top, 2.95, 3.4, border_color=a_color)

        tb = slide11.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.2), Inches(2.65), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = a_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = a_color
        p.space_after = Pt(10)

        for b in a_bullets:
            p_b = tf.add_paragraph()
            p_b.text = b
            p_b.font.size = Pt(13)
            p_b.font.color.rgb = COLOR_TEXT_BODY
            p_b.space_after = Pt(6)

    # Bottom Single Unified Verification Box
    add_card(slide11, 0.5, 5.2, 12.333, 1.85, border_color=COLOR_SUCCESS)
    tb_verif = slide11.shapes.add_textbox(Inches(0.7), Inches(5.3), Inches(11.933), Inches(1.65))
    tf_verif = tb_verif.text_frame
    tf_verif.word_wrap = True

    p_vh = tf_verif.paragraphs[0]
    p_vh.text = "FINAL VERIFICATION & CLASSIFICATION GUARANTEE"
    p_vh.font.size = Pt(14)
    p_vh.font.bold = True
    p_vh.font.color.rgb = COLOR_ACCENT_GOLD
    p_vh.space_after = Pt(4)

    p_vb = tf_verif.add_paragraph()
    p_vb.text = "• Verified Findings: Successfully validated, reproducible, security impact confirmed with generated PoC exploits.\n• Not Verified Filtered: Not Reproducible (cannot be reliably triggered) and False Positives (design is secure) are removed."
    p_vb.font.size = Pt(13)
    p_vb.font.color.rgb = COLOR_TEXT_MAIN
    p_vb.space_after = Pt(6)

    p_statement = tf_verif.add_paragraph()
    p_statement.text = '"Only unique, reproducible, and security-relevant vulnerabilities are included in the final report, ensuring high-confidence SAST validation with minimal false positives."'
    p_statement.font.size = Pt(13)
    p_statement.font.bold = True
    p_statement.font.color.rgb = COLOR_SUCCESS

    # ==========================================
    # SLIDE 12: Combined Synthesis & Remediation Roadmap
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Synthesis & Remediation Roadmap", "REMEDIATION & SUMMARY")

    add_card(slide12, 0.5, 1.65, 12.333, 5.4)
    tb = slide12.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.733), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SYNTHESIS OF DYNAMIC & STATIC SECURITY FINDINGS"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(14)

    steps = [
        ("1. Immediate High-Severity RTL Patching:", "Prioritize RTL patches for KV-002 (unconditional key zeroize), ECC-001 (verify output leak), RVCORE-001 (PMP bypass), and INTEG-001 (OTP gate hardcode)."),
        ("2. Standardize Hardware Reset Domains:", "Align lock control registers across Key Vault, PCR Vault, and Data Vault to the cold reset domain (`hard_reset_b`) to eliminate post-warm-reset secret exposure."),
        ("3. Automated Regression via Dynamic Exploits:", "Utilize the 34 developed standalone PoC exploit testbenches (`work/exploits/`) to verify that proposed RTL patches resolve vulnerabilities without introducing regressions."),
        ("4. Unified Security Sign-Off:", "Combine Phase 1 Dynamic Simulation PoC exploits and Phase 2 Python SAST & AI Validation proofs into a single consolidated audit sign-off package.")
    ]

    for title, desc in steps:
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT_BLUE
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_after = Pt(16)

    prs.save(output_path)
    print(f"Presentation updated successfully to: {output_path}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "Caliptra_RTL_Security_Analysis_Report.pptx"
    build_presentation(out_file)
